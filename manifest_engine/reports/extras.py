"""
extras.py - Reportes de control: PorBL, Chasis y PorMarca.

No reemplazan a las planillas de liquidación: son listados planos para cruzar
rápidamente qué B/L hay y cuánto factura cada uno en Buenos Aires, qué chasis
viajan y cuántos autos por marca.
"""

from collections import OrderedDict

from openpyxl import Workbook

from .. import config
from ..domain.classification import resolve_brand
from ._style import (
    BORDER_THIN, FILL_HEADER, FONT_ARIAL_10, FONT_ARIAL_10B, FONT_ARIAL_12B,
    apply_style_range, set_column_widths,
)

SIN_MARCA = "(sin marca)"


def _ops(impo_bls, expo_bls):
    """[(nombre_hoja, bls_ordenados)] para las operaciones presentes."""
    out = []
    for nombre, bls in (("IMPO", impo_bls), ("EXPO", expo_bls)):
        if bls:
            out.append((nombre, sorted(bls, key=lambda b: (b.port_of_discharge,
                                                           b.bl_no))))
    return out


def _title(ws, ship_name, voyage, op, extra=""):
    texto = f"{ship_name} {voyage} - {op}"
    if extra:
        texto += f" - {extra}"
    ws.cell(row=1, column=1, value=texto).font = FONT_ARIAL_12B


def _headers(ws, headers):
    for col, val in enumerate(headers, start=1):
        ws.cell(row=3, column=col, value=val)
    apply_style_range(ws, 3, 1, len(headers), font=FONT_ARIAL_10B,
                      border=BORDER_THIN, fill=FILL_HEADER)


def _total_buenos_aires(bl):
    if bl.op_type == "IMPO":
        return (bl.totals.get("Total Buenos Aires USD")
                or bl.totals.get("Total Buenos Aires EUR") or "")
    return bl.totals.get("Total Buenos Aires Monto") or ""


def _empty(wb, output_path):
    wb.create_sheet(title="SIN DATOS")
    wb.save(output_path)
    return output_path


# ────────────────────────────── PorBL ───────────────────────────────────

def generar_por_bl(impo_bls, expo_bls, ship_name, voyage, output_path, cfg=None):
    cfg = cfg or config.EngineConfig()
    wb = Workbook()
    wb.remove(wb.active)

    ops = _ops(impo_bls, expo_bls)
    if not ops:
        return _empty(wb, output_path)

    for op, bls in ops:
        ws = wb.create_sheet(title=op)
        set_column_widths(ws, {"A": 16, "B": 36, "C": 20, "D": 24, "E": 20})
        _title(ws, ship_name, voyage, op)
        _headers(ws, ["B/L", "CLIENTE", "PUERTO CARGA", "PUERTO DESCARGA",
                      "TOTAL BUENOS AIRES"])

        row = 4
        for bl in bls:
            ws.cell(row=row, column=1, value=bl.bl_clean)
            ws.cell(row=row, column=2, value=bl.entity)
            ws.cell(row=row, column=3, value=bl.port_of_loading)
            ws.cell(row=row, column=4, value=bl.port_of_discharge)
            ws.cell(row=row, column=5, value=_total_buenos_aires(bl))
            apply_style_range(ws, row, 1, 5, font=FONT_ARIAL_10, border=BORDER_THIN)
            row += 1

    wb.save(output_path)
    return output_path


# ────────────────────────────── Chasis ──────────────────────────────────

def _vehicle_types(bl):
    """Tipos de vehículo del B/L expandidos por cantidad ('NEW CAR', ...)."""
    tipos = []
    for tipo, cantidad in bl.cargo.vehicles.items():
        tipos.extend([tipo.upper()] * cantidad)
    return tipos


def generar_chassis(impo_bls, expo_bls, ship_name, voyage, output_path, cfg=None):
    cfg = cfg or config.EngineConfig()
    wb = Workbook()
    wb.remove(wb.active)

    ops = _ops(impo_bls, expo_bls)
    if not ops:
        return _empty(wb, output_path)

    for op, bls in ops:
        ws = wb.create_sheet(title=op)
        set_column_widths(ws, {"A": 16, "B": 24, "C": 16, "D": 24})
        _title(ws, ship_name, voyage, op, "CHASIS/VIN")
        _headers(ws, ["B/L", "MARCA", "TIPO", "VIN/CHASIS"])

        row = 4
        for bl in bls:
            if not bl.cargo.vehicles or not bl.vins:
                continue
            marca = resolve_brand(bl)
            tipos = _vehicle_types(bl)
            for i, vin in enumerate(bl.vins):
                ws.cell(row=row, column=1, value=bl.bl_clean)
                ws.cell(row=row, column=2, value=marca or "")
                ws.cell(row=row, column=3,
                        value=tipos[i] if i < len(tipos) else (tipos[-1] if tipos else ""))
                ws.cell(row=row, column=4, value=vin)
                apply_style_range(ws, row, 1, 4, font=FONT_ARIAL_10,
                                  border=BORDER_THIN)
                row += 1

    wb.save(output_path)
    return output_path


# ───────────────────────────── PorMarca ─────────────────────────────────

def generar_por_marca(impo_bls, expo_bls, ship_name, voyage, output_path, cfg=None):
    cfg = cfg or config.EngineConfig()
    wb = Workbook()
    wb.remove(wb.active)

    ops = _ops(impo_bls, expo_bls)
    if not ops:
        return _empty(wb, output_path)

    for op, bls in ops:
        ws = wb.create_sheet(title=op)
        set_column_widths(ws, {"A": 28, "B": 14, "C": 12, "D": 12})
        _title(ws, ship_name, voyage, op, "AUTOS POR MARCA")
        _headers(ws, ["MARCA", "CANT. AUTOS", "CANT. VIN", "CANT. BL"])

        resumen = OrderedDict()
        for bl in bls:
            if not bl.cargo.vehicles:
                continue
            marca = resolve_brand(bl) or SIN_MARCA
            acc = resumen.setdefault(marca, {"autos": 0, "vins": 0, "bls": 0})
            acc["autos"] += bl.cargo.n_vehicles
            acc["vins"] += len(bl.vins)
            acc["bls"] += 1

        row = 4
        for marca, acc in sorted(resumen.items(),
                                 key=lambda kv: (-kv[1]["autos"], kv[0])):
            ws.cell(row=row, column=1, value=marca)
            ws.cell(row=row, column=2, value=acc["autos"])
            ws.cell(row=row, column=3, value=acc["vins"])
            ws.cell(row=row, column=4, value=acc["bls"])
            apply_style_range(ws, row, 1, 4, font=FONT_ARIAL_10, border=BORDER_THIN)
            row += 1

        if resumen:
            ws.cell(row=row, column=1, value="TOTAL")
            for col in ("B", "C", "D"):
                ws[f"{col}{row}"] = f"=SUM({col}4:{col}{row - 1})"
            apply_style_range(ws, row, 1, 4, font=FONT_ARIAL_12B,
                              border=BORDER_THIN, fill=FILL_HEADER)

    wb.save(output_path)
    return output_path
