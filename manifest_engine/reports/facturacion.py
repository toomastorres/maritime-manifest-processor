"""
facturacion.py - 'ATL 0426 - Facturación.xlsx': lo que hay que facturar en
Buenos Aires, para enviar a Caja.

Una fila por B/L y por cada total que el manifiesto marca para pagar en Buenos
Aires (incluye 'Buenos Aires - Partner'): IMPO los collect, EXPO los prepaid. El
cliente y la CUIT son los del manifiesto (consignatario en IMPO, embarcador en
EXPO); el cliente que se factura puede ser otro, por eso hay columnas para que
Caja complete el cliente facturado y el número de factura.

La hoja DETALLE trae las líneas de cargo que forman cada total. Si el B/L tiene
otro total en la misma condición y moneda que se paga afuera (por ejemplo EXPO
con prepaid en Buenos Aires y en Casa Matriz), el manifiesto no marca cada línea:
se toma la única combinación de líneas que suma el total de Buenos Aires. Si no
es única, se listan todas y se marca que no están discriminadas.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ..rules.billing import is_local, local_lines

F_TITLE = Font(name="Arial", size=16, bold=True)
F_BOLD = Font(name="Arial", size=10, bold=True)
F_DATA = Font(name="Arial", size=10)
F_TOTAL = Font(name="Arial", size=11, bold=True)
F_NOTE = Font(name="Arial", size=9, italic=True, color="595959")
FILL_HDR = PatternFill("solid", start_color="FFD9E1F2", end_color="FFD9E1F2")
FILL_TOTAL = PatternFill("solid", start_color="FFFCE4D6", end_color="FFFCE4D6")
FILL_INPUT = PatternFill("solid", start_color="FFFFF2CC", end_color="FFFFF2CC")
FILL_WARN = PatternFill("solid", start_color="FFFFC7CE", end_color="FFFFC7CE")
_t = Side(style="thin")
_m = Side(style="medium")
B_THIN = Border(left=_t, right=_t, top=_t, bottom=_t)
B_MED = Border(left=_m, right=_m, top=_m, bottom=_m)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)
MONEY = "#,##0.00"

COLS = [("B/L", 13.3), ("CLIENTE (según manifiesto)", 38), ("CUIT", 14.5), ("PUERTO", 15),
        ("LUGAR DE PAGO", 22), ("COND.", 7), ("MONEDA", 10), ("TOTAL", 13.5),
        ("CONCEPTOS", 46), ("CLIENTE FACTURADO", 30), ("FACTURA Nº", 16)]
INPUT_COLS = (10, 11)


def generar_facturacion(impo_bls, expo_bls, ship_name, voyage, output_path):
    wb = Workbook()
    wb.remove(wb.active)
    for op, bls in (("IMPO", impo_bls), ("EXPO", expo_bls)):
        if bls is None:
            continue
        filas = _filas(bls, op)
        _hoja(wb, op, ship_name, voyage, filas)
        _detalle(wb, op, filas)
    if not wb.worksheets:
        wb.create_sheet(title="SIN DATOS")
    wb.save(output_path)
    return output_path


def _filas(bls, op):
    """(bl, total, líneas del total, discriminado, otros) por cada total pagadero en Buenos Aires."""
    out = []
    for bl in sorted(bls, key=lambda b: b.bl_clean):
        for t in bl.total_lines:
            if not is_local(t.place):
                continue
            otros = [o for o in bl.total_lines if o is not t and o.letra == t.letra and o.moneda == t.moneda]
            lineas = local_lines(bl, t.letra, t.moneda)
            if lineas is None:
                lineas = [cl for cl in bl.charge_lines if cl.letra == t.letra and cl.moneda == t.moneda]
                out.append((bl, t, lineas, False, otros))
            else:
                out.append((bl, t, lineas, True, otros))
    return out


def _hoja(wb, op, ship, voyage, filas):
    ws = wb.create_sheet(title=op)
    quien = "consignatario" if op == "IMPO" else "embarcador"
    ws["A1"] = f"A FACTURAR EN BUENOS AIRES — {op} — {ship} {voyage}"
    ws["A1"].font = F_TITLE
    ws.row_dimensions[1].height = 20.25
    ws["A2"] = (f"Cliente y CUIT según el manifiesto ({quien}). El cliente facturado puede ser otro: "
                f"completar las columnas en amarillo.")
    ws["A2"].font = F_NOTE
    for i, (t, w) in enumerate(COLS, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
        c = ws.cell(row=4, column=i, value=t)
        c.font = F_BOLD
        c.fill = FILL_INPUT if i in INPUT_COLS else FILL_HDR
        c.border = B_THIN
        c.alignment = CENTER
    ws.row_dimensions[4].height = 27

    r = 5
    for bl, t, lineas, discriminado, otros in filas:
        puerto = bl.port_of_loading if op == "IMPO" else bl.port_of_discharge
        if discriminado:
            conceptos = ", ".join(dict.fromkeys(cl.concept for cl in lineas))
        else:
            conceptos = ("No discriminado en el manifiesto: también hay "
                         + ", ".join(f"Total[{o.place}] {o.amount:,.2f}{o.letra} {o.moneda}" for o in otros)
                         + " (ver DETALLE)")
        vals = [bl.bl_clean, bl.entity if bl.entity not in ("", "Nulo") else None, bl.party_id or None,
                puerto, t.place, t.letra, t.moneda, t.amount, conceptos, None, None]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = F_DATA
            c.border = B_THIN
            if i == 8:
                c.number_format = MONEY
            if i in INPUT_COLS:
                c.fill = FILL_INPUT
            if i == 9:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if not discriminado:
                    c.fill = FILL_WARN
        r += 1
    ultima = r - 1

    # Totales por moneda.
    r += 1
    monedas = sorted({t.moneda for _, t, *_ in filas})
    for moneda in monedas:
        ws.cell(row=r, column=6, value="TOTAL")
        ws.cell(row=r, column=7, value=moneda)
        ws.cell(row=r, column=8, value=f'=SUMIF(G5:G{max(ultima, 5)},"{moneda}",H5:H{max(ultima, 5)})')
        for i in (6, 7, 8):
            c = ws.cell(row=r, column=i)
            c.font = F_TOTAL
            c.fill = FILL_TOTAL
            c.border = B_MED
        ws.cell(row=r, column=8).number_format = MONEY
        r += 1
    if not filas:
        ws.cell(row=5, column=1, value="El manifiesto no trae totales a pagar en Buenos Aires.").font = F_NOTE

    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:K{max(ultima, 5)}"
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _detalle(wb, op, filas):
    ws = wb.create_sheet(title=f"DETALLE {op}")
    cols = [("B/L", 13.3), ("LUGAR DE PAGO", 22), ("CONCEPTO", 30), ("FACTOR", 10), ("BASE", 7),
            ("TARIFA", 10), ("IMPORTE", 13), ("COND.", 7), ("MONEDA", 8.5), ("OBSERVACIÓN", 50)]
    for i, (t, w) in enumerate(cols, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
        c = ws.cell(row=1, column=i, value=t)
        c.font = F_BOLD
        c.fill = FILL_HDR
        c.border = B_THIN
        c.alignment = CENTER
    r = 2
    for bl, t, lineas, discriminado, otros in filas:
        obs = None if discriminado else "No discriminado: puede ser de otro lugar de pago"
        for cl in lineas:
            vals = [bl.bl_clean, t.place, cl.concept, cl.factor, cl.basis, cl.rate, cl.total,
                    cl.letra, cl.moneda, obs]
            for i, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=i, value=v)
                c.font = F_DATA
                c.border = B_THIN
                if i == 7:
                    c.number_format = MONEY
                if i == 10 and obs:
                    c.fill = FILL_WARN
            r += 1
        # Total del manifiesto, para controlar contra la suma de las líneas.
        for i, v in ((3, f"Total[{t.place}]"), (7, t.amount), (8, t.letra), (9, t.moneda)):
            c = ws.cell(row=r, column=i, value=v)
            c.font = F_BOLD
            c.fill = FILL_TOTAL
            c.border = B_THIN
            if i == 7:
                c.number_format = MONEY
        ws.cell(row=r, column=1, value=bl.bl_clean).font = F_BOLD
        r += 2
    ws.freeze_panes = "A2"
