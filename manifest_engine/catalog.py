"""
catalog.py - Reglas de negocio gobernadas por usuario en una planilla Excel.

El motor NO clasifica por su cuenta los conceptos de cargo, los tipos de bulto
ni los puertos: los lee de `Definiciones_reglas_comisiones.xlsx`, que completa
el equipo operativo. Lo que no está en esa planilla (o está como 'NO SÉ') se
trata como indefinido y el B/L afectado va a revisión humana.

Hojas que se leen:
  CONCEPTOS DE CARGO  A=concepto  F=qué es para la comisión  G=suma/resta
  TIPOS DE BULTO      A=tipo      G=familia  H=mínimo por contenedor  I=Control Equipment Fees
  PUERTOS EXPO        A=puerto    B=país     D=¿tasa Brasil?
"""

import os
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Optional

import openpyxl

from . import config

DEFAULT_FILENAME = "Definiciones_reglas_comisiones.xlsx"
ENV_VAR = "MANIFEST_DEFINICIONES"

# Categorías de concepto (hoja CONCEPTOS DE CARGO, columna F).
FLETE = "FLETE"
RECARGO = "RECARGO"        # incluido en el flete: si está, el B/L NO es ALL IN
GASTO = "GASTO"            # fuera del flete: no afecta ALL IN
DESCUENTO = "DESCUENTO"

# Efecto sobre el monto comisionable (columna G).
SUMA = "SUMA"
RESTA = "RESTA"
NO_AFECTA = "NO AFECTA"

# Familias de carga (hoja TIPOS DE BULTO, columna G).
CONTENEDOR = "CONTENEDOR"
BBULK = "BBULK"
CAR = "CAR"
VAN = "VAN"
RORO = "RORO"
RODADOS = (CAR, VAN, RORO)


class CatalogError(Exception):
    pass


@dataclass(frozen=True)
class ConceptRule:
    name: str
    category: Optional[str]     # FLETE | RECARGO | GASTO | DESCUENTO | None (indefinido)
    effect: Optional[str]       # SUMA | RESTA | NO AFECTA | None (indefinido)

    @property
    def defined(self) -> bool:
        return self.category is not None and self.effect is not None


@dataclass(frozen=True)
class KindRule:
    name: str
    family: Optional[str]       # CONTENEDOR | BBULK | CAR | VAN | RORO | None
    min_ctr: Optional[bool]     # cuenta para el mínimo en USD por contenedor
    cef: Optional[bool]         # cuenta para Control Equipment Fees

    @property
    def is_equipment(self) -> bool:
        """Mafi: es el equipo que lleva la carga Bbulk, no la carga en sí."""
        return any(tok.lower() in self.name.lower() for tok in config.EQUIPMENT_KIND_TOKENS)

    @property
    def defined(self) -> bool:
        if self.family is None:
            return False
        if self.family == CONTENEDOR:
            return self.min_ctr is not None and self.cef is not None
        return True


@dataclass(frozen=True)
class PortRule:
    name: str
    country: str
    brasil: Optional[bool]


def _category(v) -> Optional[str]:
    s = str(v or "").strip().upper()
    if s.startswith("FLETE"):
        return FLETE
    if s.startswith("RECARGO"):
        return RECARGO
    if s.startswith("GASTO"):
        return GASTO
    if s.startswith("DESCUENTO"):
        return DESCUENTO
    return None


def _effect(v) -> Optional[str]:
    s = str(v or "").strip().upper()
    return s if s in (SUMA, RESTA, NO_AFECTA) else None


def _family(v) -> Optional[str]:
    s = str(v or "").strip().upper()
    return s if s in (CONTENEDOR, BBULK, CAR, VAN, RORO) else None


def _yesno(v) -> Optional[bool]:
    s = str(v or "").strip().upper()
    if s in ("SI", "SÍ"):
        return True
    if s == "NO":
        return False
    return None


def _norm(s) -> str:
    return " ".join(str(s or "").split())


class Catalog:
    def __init__(self, path: Path):
        self.path = Path(path)
        self.concepts: Dict[str, ConceptRule] = {}
        self.kinds: Dict[str, KindRule] = {}
        self.ports: Dict[str, PortRule] = {}
        self._load()

    def _load(self):
        if not self.path.exists():
            raise CatalogError(f"No encuentro la planilla de definiciones: {self.path}")
        wb = openpyxl.load_workbook(self.path, data_only=True)
        for hoja in ("CONCEPTOS DE CARGO", "TIPOS DE BULTO", "PUERTOS EXPO"):
            if hoja not in wb.sheetnames:
                raise CatalogError(f"La planilla {self.path.name} no tiene la hoja '{hoja}'")

        ws = wb["CONCEPTOS DE CARGO"]
        for r in range(5, ws.max_row + 1):
            name = _norm(ws.cell(r, 1).value)
            # Fila de datos: tiene nombre y algo en ejemplo/clasificación (así se
            # pueden agregar conceptos nuevos a mano sin completar el ejemplo).
            if not name or name.startswith("Pendientes"):
                continue
            if all(ws.cell(r, c).value in (None, "") for c in (4, 6, 7)):
                continue
            self.concepts[name] = ConceptRule(name, _category(ws.cell(r, 6).value),
                                              _effect(ws.cell(r, 7).value))

        ws = wb["TIPOS DE BULTO"]
        for r in range(4, ws.max_row + 1):
            name = _norm(ws.cell(r, 1).value)
            if not name or name.startswith("Pendientes"):
                continue
            if all(ws.cell(r, c).value in (None, "") for c in (6, 7)):
                continue
            self.kinds[name] = KindRule(name, _family(ws.cell(r, 7).value),
                                        _yesno(ws.cell(r, 8).value), _yesno(ws.cell(r, 9).value))

        ws = wb["PUERTOS EXPO"]
        for r in range(5, ws.max_row + 1):
            name = _norm(ws.cell(r, 1).value)
            if not name or name.startswith("Pendientes"):
                continue
            if all(ws.cell(r, c).value in (None, "") for c in (3, 4)):
                continue
            self.ports[name] = PortRule(name, _norm(ws.cell(r, 2).value), _yesno(ws.cell(r, 4).value))

        if not self.concepts or not self.kinds:
            raise CatalogError(f"La planilla {self.path.name} no tiene conceptos o tipos de bulto cargados")

    def concept(self, name: str) -> Optional[ConceptRule]:
        return self.concepts.get(_norm(name))

    def kind(self, name: str) -> Optional[KindRule]:
        return self.kinds.get(_norm(name))

    def port(self, name: str) -> Optional[PortRule]:
        return self.ports.get(_norm(name))


def default_path() -> Path:
    """Planilla de definiciones: variable de entorno, raíz del proyecto o carpeta actual."""
    env = os.environ.get(ENV_VAR)
    if env:
        return Path(env)
    for base in (Path(__file__).resolve().parent.parent, Path.cwd()):
        p = base / DEFAULT_FILENAME
        if p.exists():
            return p
    return Path(__file__).resolve().parent.parent / DEFAULT_FILENAME


_cache: Dict[Path, Catalog] = {}


def load_catalog(path=None) -> Catalog:
    p = Path(path) if path else default_path()
    p = p.resolve()
    if p not in _cache:
        _cache[p] = Catalog(p)
    return _cache[p]


def register_new_concepts(found: Dict[str, dict], origen: str, path=None):
    """
    Agrega a la hoja CONCEPTOS DE CARGO los conceptos que aparecieron en un
    manifiesto y no estaban en la lista. Se cargan como 'no afecta' (no influyen
    en la comisión) y quedan marcados para que el equipo los revise.

    found: {concepto: {"IMPO": n_bl, "EXPO": n_bl, "ejemplo": texto}}
    Devuelve (agregados, error). Si la planilla está abierta en Excel no se
    puede guardar: se devuelve el error y no se pierde nada (el aviso queda en
    la planilla de comisiones y en pantalla).
    """
    from copy import copy
    from datetime import date

    p = Path(path) if path else default_path()
    try:
        wb = openpyxl.load_workbook(p)
        ws = wb["CONCEPTOS DE CARGO"]
    except Exception as e:  # noqa: BLE001
        return [], f"No se pudo abrir {p.name}: {e}"

    existentes = {_norm(ws.cell(r, 1).value) for r in range(5, ws.max_row + 1)}
    nuevos = {k: v for k, v in found.items() if _norm(k) not in existentes}
    if not nuevos:
        return [], None

    # Última fila de conceptos y fila del contador 'Pendientes de responder'.
    last = 4
    pend_row = None
    for r in range(5, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.startswith("Pendientes"):
            pend_row = r
        elif a not in (None, ""):
            last = r
    if pend_row:
        for c in (1, 2):
            ws.cell(pend_row, c).value = None

    modelo = last
    hoy = date.today().strftime("%d/%m/%Y")
    r = last
    for nombre, d in sorted(nuevos.items()):
        r += 1
        valores = [
            nombre, d.get("IMPO") or None, d.get("EXPO") or None, d.get("ejemplo"),
            "Nuevo: no afecta la comisión", "GASTO FUERA DEL FLETE (no afecta ALL IN)",
            "NO AFECTA", f"Agregado por el motor el {hoy} ({origen}). Revisar clasificación.",
        ]
        for c, v in enumerate(valores, start=1):
            cell = ws.cell(r, c, v)
            src = ws.cell(modelo, c)
            if src.has_style:
                cell.font = copy(src.font)
                cell.border = copy(src.border)
                cell.alignment = copy(src.alignment)
        for dv in ws.data_validations.dataValidation:
            f1 = str(dv.formula1 or "")
            if "FLETE BASE" in f1:
                dv.add(f"F{r}")
            elif "RESTA" in f1:
                dv.add(f"G{r}")

    if pend_row:
        ws.cell(r + 2, 1, "Pendientes de responder:").font = copy(ws.cell(5, 1).font)
        ws.cell(r + 2, 2, f"=COUNTBLANK(F5:F{r})")
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A4:H{r}"

    try:
        wb.save(p)
    except PermissionError:
        return [], (f"No se pudo guardar {p.name} (¿está abierta en Excel?). "
                    f"Conceptos nuevos sin registrar: {', '.join(sorted(nuevos))}")
    _cache.pop(p.resolve(), None)
    return sorted(nuevos), None
