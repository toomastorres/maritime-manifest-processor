"""
Suite de pruebas unitarias e integrales del motor de manifiestos marítimos v2.0.

Cubre:
  1. Catálogo gobernado en Excel (`catalog.py`) y auto-registro idempotente de
     conceptos nuevos sobre copia temporal (`register_new_concepts`).
  2. Parser IMPO/EXPO (`parsing/manifest.py`): descarte de equipos vacíos `P3`,
     encabezados de cargo partidos en dos líneas, múltiples puertos y equipos Mafi.
  3. Clasificación de dominio (`domain/classification.py`): VIN, marca y sufijo forwarder.
  4. Validación Módulo 11 de CUIT y liquidación de Caja (`rules/billing.py`):
     subagente `Buenos Aires - Partner`, búsqueda combinatoria de líneas prepagas
     locales y control de cuadratura `Total[...]`.
  5. Motor de comisiones (`rules/commission.py`): condición `ALL IN` (80 %),
     excepción asimétrica `THC (POL)` en IMPO, mínimo por contenedor, tasa
     regional Brasil, modos de rodados (`AMA`, `UNIDAD`, `ESTIMADO`) y envío a
     `REVISIÓN` ante moneda no convertible o puerto desconocido.
  6. Partición determinística de cargas mixtas (`rules/split.py`): separación
     exacta por medidas vs bloqueo sin inventar importes ante flete global (`AA`).
  7. Auditoría híbrida de gastos de línea (`validation.py`) y generación
     end-to-end de las 9 planillas Excel del viaje + plantillas modelo (`--modelos`).
"""

import shutil
from copy import deepcopy
from pathlib import Path

import openpyxl
import pytest

from manifest_engine import config
from manifest_engine.__main__ import parse_agency, process_directory
from manifest_engine.catalog import (
    BBULK,
    CAR,
    CONTENEDOR,
    DESCUENTO,
    FLETE,
    GASTO,
    NO_AFECTA,
    RECARGO,
    RESTA,
    RORO,
    SUMA,
    VAN,
    default_path,
    load_catalog,
    register_new_concepts,
)
from manifest_engine.domain import agency_suffix, detect_make, format_cargo, resolve_brand
from manifest_engine.domain.models import BLRecord, CargoItem, ChargeLine, TotalLine
from manifest_engine.parsing.manifest import cuit_ok, extract_party_id, parse_manifest
from manifest_engine.reports.comisiones import generar_modelo_comisiones
from manifest_engine.rules.billing import expo_money, impo_money, local_lines, total_mismatches
from manifest_engine.rules.commission import (
    MODE_AGENCY,
    MODE_ESTIMADO,
    MODE_UNIT,
    evaluate,
)
from manifest_engine.rules.split import MIXTO, filas_del_bl, split_alerts
from manifest_engine.validation import revisar_bls

REPO_ROOT = Path(__file__).resolve().parents[1]
SAMPLES = REPO_ROOT / "samples"
IMPO_CSV = str(SAMPLES / "ATL0426_IMPO_1.CSV")
EXPO_CSV = str(SAMPLES / "ATL0426_EXPO_1.CSV")


@pytest.fixture(scope="module")
def catalog():
    return load_catalog()


@pytest.fixture(scope="module")
def impo(catalog):
    return parse_manifest(IMPO_CSV, "IMPO", catalog=catalog)


@pytest.fixture(scope="module")
def expo(catalog):
    return parse_manifest(EXPO_CSV, "EXPO", catalog=catalog)


# ── 1. Catálogo gobernado en Excel ────────────────────────────────────────

def test_catalog_loads_all_rules_defined(catalog):
    assert len(catalog.concepts) >= 80
    assert len(catalog.kinds) >= 35
    assert len(catalog.ports) >= 18
    assert all(r.defined for r in catalog.concepts.values())
    assert all(r.defined for r in catalog.kinds.values())

    assert catalog.concept("Basic Frt.").category == FLETE
    assert catalog.concept("Basic Frt.").effect == SUMA
    assert catalog.concept("BAF").category == RECARGO
    assert catalog.concept("Toll S/C(POD)").category == GASTO
    assert catalog.concept("FAC (x% of y)").category == DESCUENTO
    assert catalog.concept("FAC (x% of y)").effect == RESTA

    ctr = catalog.kind("40 ft. High Cube")
    assert ctr.family == CONTENEDOR and ctr.min_ctr is True and ctr.cef is True

    mafi = catalog.kind("Mafi 40")
    assert mafi.family == BBULK and mafi.is_equipment is True
    assert mafi.min_ctr is False and mafi.cef is False

    assert catalog.port("SANTOS").brasil is True
    assert catalog.port("HAMBURG").brasil is False


def test_catalog_register_new_concepts_on_temp_copy(tmp_path):
    copy_xlsx = tmp_path / "Definiciones_reglas_comisiones.xlsx"
    shutil.copy2(default_path(), copy_xlsx)

    found = {
        "Quantum Ice S/C": {
            "IMPO": 2,
            "EXPO": 0,
            "ejemplo": "S300000199: Quantum Ice S/C 1 PU = 75.00P USD",
        }
    }
    added, err = register_new_concepts(found, "ATL 0426", path=copy_xlsx)
    assert err is None
    assert added == ["Quantum Ice S/C"]

    reloaded = load_catalog(copy_xlsx)
    rule = reloaded.concept("Quantum Ice S/C")
    assert rule is not None and rule.defined is True
    assert rule.category == GASTO and rule.effect == NO_AFECTA

    # Idempotencia: volver a registrar el mismo concepto no duplica filas.
    added_again, err_again = register_new_concepts(found, "ATL 0426", path=copy_xlsx)
    assert err_again is None and added_again == []


# ── 2. Parsing IMPO/EXPO y clasificación de dominio ───────────────────────

def test_impo_discards_p3_and_parses_s3_bls(impo):
    bl_nos = [b.bl_no for b in impo]
    assert "P300000901" not in bl_nos
    assert bl_nos == [
        "S300000101",
        "S300000102",
        "S300000103",
        "S300000104",
        "S300000105",
    ]
    assert all(b.port_of_discharge == "ZARATE" for b in impo)


def test_expo_multi_port_and_p3_discard(expo):
    bl_nos = [b.bl_no for b in expo]
    assert "P300000902" not in bl_nos
    assert bl_nos == ["S300000201", "S300000202", "S300000203", "S300000204"]
    ports = {b.bl_no: b.port_of_discharge for b in expo}
    assert ports["S300000201"] == "HAMBURG"
    assert ports["S300000202"] == "HAMBURG"
    assert ports["S300000203"] == "SANTOS"
    assert ports["S300000204"] == "SANTOS"


def test_vin_brand_and_agency_suffix(impo):
    bl_cars = impo[2]  # S300000103
    assert len(bl_cars.vins) == 5
    assert bl_cars.vins[0] == "ZZDLTR8K9AA100001"
    assert detect_make(bl_cars.description) == "DELTAR"
    assert agency_suffix(bl_cars.entity) == "OCEANLINK"
    assert resolve_brand(bl_cars) == "DELTAR - OCEANLINK"
    assert format_cargo(bl_cars) == "5 NEW CAR - DELTAR - OCEANLINK"

    bl_vans = impo[3]  # S300000104
    assert bl_vans.units(VAN) == 4
    assert format_cargo(bl_vans) == "4 NEW VAN - CORVEX"


def test_wrapped_two_line_charge_headers_resolved(impo):
    # En S300000101 'Toll' + 'S/C(POD)' viene partido en 2 líneas
    concepts_101 = [cl.concept for cl in impo[0].charge_lines]
    assert "Toll S/C(POD)" in concepts_101
    # En S300000105 'ETS+FUEL EU' + 'Surcharge' viene partido en 2 líneas
    concepts_105 = [cl.concept for cl in impo[4].charge_lines]
    assert "ETS+FUEL EU Surcharge" in concepts_105


def test_mafi_equipment_excluded_when_carrying_bbulk(expo):
    bl_mafi = expo[3]  # S300000204: 1-Mafi 40 + 2-CASE(S)
    assert len(bl_mafi.items) == 2
    assert len(bl_mafi.cargo_items) == 1
    assert bl_mafi.cargo_items[0].kind == "CASE(S)"
    assert bl_mafi.families() == [BBULK]
    assert format_cargo(bl_mafi) == "2 CASE"


# ── 3. CUIT Módulo 11 y Facturación para Caja ─────────────────────────────

def test_cuit_modulo_11_and_party_extraction(impo, expo):
    assert cuit_ok("30710001010") is True
    assert cuit_ok("30710001019") is False  # dígito verificador alterado
    assert cuit_ok("123") is False

    # IMPO toma CUIT del Consignee (CN), EXPO del Shipper (SH), incluso partido en 2 líneas
    assert impo[0].party_id == "30-71000101-0"
    assert expo[0].party_id == "30-72000201-5"

    split_lines = [
        "SH:EXPORTADORA DEMO SA",
        "AV. PUERTO 500 - C.U.I.T.",
        "30-71000101-0 BUENOS AIRES",
        "CN:FOREIGN BUYER INC",
        "CUIT 30-71000102-9",
    ]
    assert extract_party_id(split_lines, "SH") == "30-71000101-0"
    assert extract_party_id(split_lines, "CN") == "30-71000102-9"


def test_impo_billing_and_partner_subagent(impo):
    # S300000101: Total[Buenos Aires] 846C USD = THC 360 + Toll 450 + Sweeping 36 -> neto collect 0
    m1 = impo_money(impo[0])
    assert m1.local_total == {"USD": 846.0}
    assert m1.collect == {}
    assert m1.warnings == []

    # S300000102: Total[Buenos Aires - Partner] 250C USD reconocido como cobro local
    m2 = impo_money(impo[1])
    assert m2.local_total == {"USD": 250.0}
    assert m2.warnings == []

    # S300000103: Rodados con THC por unidad (450 USD) sin columna propia -> queda en collect USD
    m3 = impo_money(impo[2])
    assert m3.local_total == {"USD": 600.0}
    assert m3.collect == {"USD": 450.0}


def test_expo_combinatorial_local_prepaid_matching(expo):
    bl_201 = expo[0]  # 6 líneas P USD (3.135 USD) repartidas en Buenos Aires (585) + Head Office (2.550)
    em = expo_money(bl_201)
    assert em.local_prepaid == pytest.approx(585.0)
    assert em.abroad == {"USD": 2550.0}
    assert em.collect == {"EUR": 24.0}

    matched = local_lines(bl_201, "P", "USD")
    assert matched is not None
    assert sorted(cl.concept for cl in matched) == [
        "ENS",
        "Sweeping",
        "THC (POL)",
        "Toll S/C(POL)",
    ]
    assert sum(cl.total for cl in matched) == pytest.approx(585.0)

    # Si dos conceptos distintos empatan en importe (ej. BAF = 300 y Toll = 300),
    # la combinación deja de ser única y el motor devuelve None sin adivinar.
    ambiguous_bl = deepcopy(bl_201)
    for cl in ambiguous_bl.charge_lines:
        if cl.concept == "BAF":
            cl.total = 300.0
    assert local_lines(ambiguous_bl, "P", "USD") is None


def test_total_mismatches_control(impo, expo):
    for b in impo + expo:
        assert total_mismatches(b) == [], f"Descuadre inesperado en {b.bl_no}"

    tampered = deepcopy(impo[0])
    tampered.total_lines[0].amount += 50.0
    mismatches = total_mismatches(tampered)
    assert len(mismatches) == 1
    assert mismatches[0][:2] == ("P", "USD")


# ── 4. Motor de comisiones gobernado por catálogo ─────────────────────────

def test_commission_container_minimum_and_cef(impo, catalog):
    cfg = config.EngineConfig(roe=1.08, agency_bls=frozenset())
    res = evaluate(impo[0], "IMPO", cfg, catalog=catalog)  # S300000101: 3x40', Basic 1800 USD + BAF
    assert res.ok is True
    assert res.all_in is False
    assert res.cef_units == 3
    part = res.parts[0]
    assert part.calc_usd == pytest.approx(36.0)     # 1800 * 2%
    assert part.minimum_usd == pytest.approx(75.0)  # 3 * 25 USD
    assert res.total_usd == pytest.approx(75.0)


def test_commission_all_in_80_percent_and_brasil_rate(impo, expo, catalog):
    cfg = config.EngineConfig(roe=1.08, agency_bls=frozenset())

    # S300000102: IMPO ALL IN (2000 * 80% = 1600 * 2% = 32 USD > mín 25 USD)
    r_impo = evaluate(impo[1], "IMPO", cfg, catalog=catalog)
    assert r_impo.ok and r_impo.all_in is True and r_impo.brasil is False
    assert r_impo.parts[0].base == pytest.approx(1600.0)
    assert r_impo.total_usd == pytest.approx(32.0)

    # S300000203: EXPO SANTOS (Brasil 3,5%) + ALL IN (2400 * 80% = 1920 * 3,5% = 67,20 USD > mín 50 USD)
    r_expo = evaluate(expo[2], "EXPO", cfg, catalog=catalog)
    assert r_expo.ok and r_expo.all_in is True and r_expo.brasil is True
    assert r_expo.parts[0].base == pytest.approx(1920.0)
    assert r_expo.total_usd == pytest.approx(67.20)


def test_impo_asymmetric_thc_pol_breaks_all_in(catalog):
    bl = BLRecord(bl_no="S300000901", op_type="IMPO", port_of_loading="ANTWERP", port_of_discharge="ZARATE")
    bl.items.append(CargoItem(qty=1, kind="20 ft. Dry Cargo", family=CONTENEDOR))
    bl.rebuild_cargo()
    bl.charge_lines = [
        ChargeLine(desc="basic frt", total=2000.0, letra="P", moneda="USD", concept="Basic Frt."),
        ChargeLine(desc="thc (pol)", total=100.0, letra="P", moneda="EUR", concept="THC (POL)"),
    ]
    res = evaluate(bl, "IMPO", config.EngineConfig(), catalog=catalog)
    assert res.all_in is False
    assert "THC (POL)" in res.all_in_breakers


def test_rolling_modes_ama_vs_unit_vs_estimado(impo, catalog):
    bl_cars = impo[2]  # S300000103: 5 New Car(s), Basic Frt 2000 EUR, BAF 200 EUR (no ALL IN)

    # Modo AMA (booking cerrado por la agencia): 2000 EUR * 1.08 ROE * 1% = 21.60 USD
    res_ama = evaluate(bl_cars, "IMPO", config.EngineConfig(roe=1.08, agency_bls={"S300000103"}), catalog)
    assert res_ama.rodado_mode == MODE_AGENCY
    assert res_ama.total_usd == pytest.approx(21.60)

    # Modo UNIDAD (lista confirmada sin este B/L): 5 cars * 2.50 USD = 12.50 USD
    res_unit = evaluate(bl_cars, "IMPO", config.EngineConfig(roe=1.08, agency_bls=frozenset()), catalog)
    assert res_unit.rodado_mode == MODE_UNIT
    assert res_unit.total_usd == pytest.approx(12.50)

    # Modo ESTIMADO (sin lista de Comercial): mismo importe por unidad pero marcado ESTIMADO
    res_est = evaluate(bl_cars, "IMPO", config.EngineConfig(roe=1.08, agency_bls=None), catalog)
    assert res_est.rodado_mode == MODE_ESTIMADO
    assert res_est.total_usd == pytest.approx(12.50)


def test_never_invent_values_on_unconvertible_currency_or_unknown_port(catalog):
    # Moneda GBP no convertible -> REVISION con total_usd = None
    bl_gbp = BLRecord(bl_no="S300000902", op_type="IMPO", port_of_loading="TILBURY", port_of_discharge="ZARATE")
    bl_gbp.items.append(CargoItem(qty=1, kind="20 ft. Dry Cargo", family=CONTENEDOR))
    bl_gbp.rebuild_cargo()
    bl_gbp.charge_lines = [
        ChargeLine(desc="basic frt", total=1500.0, letra="P", moneda="GBP", concept="Basic Frt."),
    ]
    r_gbp = evaluate(bl_gbp, "IMPO", config.EngineConfig(), catalog)
    assert r_gbp.ok is False and r_gbp.status == "REVISION" and r_gbp.total_usd is None
    assert any("GBP" in msg for msg in r_gbp.reasons)

    # Puerto EXPO no registrado en catálogo -> REVISION con total_usd = None
    bl_port = BLRecord(bl_no="S300000903", op_type="EXPO", port_of_loading="ZARATE", port_of_discharge="VALPARAISO")
    bl_port.items.append(CargoItem(qty=1, kind="20 ft. Dry Cargo", family=CONTENEDOR))
    bl_port.rebuild_cargo()
    bl_port.charge_lines = [
        ChargeLine(desc="basic frt", total=1500.0, letra="P", moneda="USD", concept="Basic Frt."),
    ]
    r_port = evaluate(bl_port, "EXPO", config.EngineConfig(), catalog)
    assert r_port.ok is False and r_port.total_usd is None
    assert any("VALPARAISO" in msg for msg in r_port.reasons)


# ── 5. Partición determinística de cargas mixtas (`split.py`) ─────────────

def test_mixed_cargo_exact_split_freight_and_toll(impo, catalog):
    bl_mix = impo[4]  # S300000105: 1 New LM RoRo (6.5 LM, 18 Tn) + 4 CASE(S) (32 CBM, 5 Tn)
    cfg = config.EngineConfig(roe=1.08, agency_bls=frozenset())
    res = evaluate(bl_mix, "IMPO", cfg, catalog)
    assert res.ok is True and res.split_ok is True
    # 1 RoRo * 12.0 USD + (2560 EUR * 1.08 * 2% = 55.296 USD) = 67.296 USD
    assert res.total_usd == pytest.approx(12.0 + 2560.0 * 1.08 * 0.02)

    rows = filas_del_bl(bl_mix, res, "IMPO")
    assert split_alerts(rows) == []
    assert len(rows) == 2

    row_roro, row_bbulk = rows[0], rows[1]
    assert row_roro.tipo == "RODADO"
    assert row_roro.qty(RORO) == 1 and row_roro.qty(BBULK) == 0
    assert sum(cl.total for cl in row_roro.comm_lines) == pytest.approx(2600.0)
    assert row_roro.toll == pytest.approx(360.0)

    assert row_bbulk.tipo == "BBULK"
    assert row_bbulk.qty(BBULK) == 4 and row_bbulk.qty(RORO) == 0
    assert sum(cl.total for cl in row_bbulk.comm_lines) == pytest.approx(2560.0)
    assert row_bbulk.toll == pytest.approx(100.0)


def test_mixed_cargo_ambiguous_lumpsum_stays_mixto_and_alerts(catalog):
    bl_amb = BLRecord(bl_no="S300000904", op_type="IMPO", port_of_loading="ANTWERP", port_of_discharge="ZARATE")
    bl_amb.items = [
        CargoItem(qty=1, kind="New LM RoRo", family=RORO, tn=12.0, lm=6.0),
        CargoItem(qty=2, kind="CASE(S)", family=BBULK, tn=4.0, cbm=10.0),
    ]
    bl_amb.rebuild_cargo()
    cl_frt = ChargeLine(
        desc="basic frt", total=3000.0, letra="P", moneda="USD",
        factor=1.0, basis="AA", rate=3000.0, concept="Basic Frt.",
    )
    bl_amb.charge_lines = [cl_frt]
    bl_amb.set_charge("Basic FRT", 3000.0, "P", "USD")

    # Con booking AMA en rodados, ambas cargas requieren repartir el flete: '1 AA' es ambiguo.
    res = evaluate(bl_amb, "IMPO", config.EngineConfig(agency_bls={"S300000904"}), catalog)
    assert res.ok is False and res.split_ok is False and res.total_usd is None

    rows = filas_del_bl(bl_amb, res, "IMPO")
    assert len(rows) == 1 and rows[0].tipo == MIXTO
    alerts = split_alerts(rows)
    assert len(alerts) == 1 and "MIXTO" in alerts[0]


# ── 6. Auditoría REVISIÓN, parser CLI y generación end-to-end ─────────────

def test_parse_agency_cli_helper():
    assert parse_agency("") == (None, [])
    assert parse_agency("ninguno") == (frozenset(), [])
    valid, invalid = parse_agency("S300000103, S300000202[T] INVALID99")
    assert valid == frozenset({"S300000103", "S300000202"})
    assert invalid == ["INVALID99"]


def test_hybrid_line_cost_validation_flags_sweeping_typo(impo, expo, catalog):
    cfg = config.EngineConfig(roe=1.08, agency_bls=frozenset({"S300000103"}))
    impo_alerts = revisar_bls(impo, "IMPO", cfg, catalog=catalog)
    assert len(impo_alerts) == 1
    assert impo_alerts[0].bl == "S300000101"
    assert impo_alerts[0].tipo == "GASTO LINEA"
    assert "SWEEPING" in impo_alerts[0].detalle

    expo_alerts = revisar_bls(expo, "EXPO", cfg, catalog=catalog)
    assert expo_alerts == []


def test_end_to_end_process_directory_generates_9_workbooks(tmp_path):
    process_directory(
        str(SAMPLES),
        str(tmp_path),
        roe=1.08,
        agency_override=frozenset({"S300000103"}),
        register_concepts=False,
    )
    voyage_dir = tmp_path / "ATL 0426"
    assert voyage_dir.is_dir()
    assert (voyage_dir / "Manifiestos" / "ATL0426_IMPO_1.CSV").is_file()
    assert (voyage_dir / "Manifiestos" / "ATL0426_EXPO_1.CSV").is_file()

    reports = sorted(p.name for p in voyage_dir.glob("*.xlsx"))
    assert len(reports) == 9
    expected_suffixes = (
        "IMPO.xlsx",
        "Comm. Inward.xlsx",
        "EXPO.xlsx",
        "Comm. Outward.xlsx",
        "ENS.xlsx",
        "Facturación.xlsx",
        "PorBL.xlsx",
        "Chasis.xlsx",
        "PorMarca.xlsx",
    )
    for suffix in expected_suffixes:
        assert any(r.endswith(suffix) for r in reports), f"Falta reporte {suffix}"

    wb_fact = openpyxl.load_workbook(voyage_dir / "ATL 0426 - Facturación.xlsx")
    assert {"IMPO", "EXPO", "DETALLE IMPO", "DETALLE EXPO"} <= set(wb_fact.sheetnames)

    wb_comm = openpyxl.load_workbook(voyage_dir / "ATL 0426 - Comm. Inward.xlsx")
    assert {"Comm Inward", "DETALLE B-L"} <= set(wb_comm.sheetnames)


def test_generate_empty_commission_templates(tmp_path):
    expected_sheet = {"IMPO": "Comm Inward", "EXPO": "Comm Outward"}
    for op in ("IMPO", "EXPO"):
        out = tmp_path / f"MODELO_{op}.xlsx"
        generar_modelo_comisiones(op, str(out))
        assert out.is_file()
        wb = openpyxl.load_workbook(out)
        assert expected_sheet[op] in wb.sheetnames
