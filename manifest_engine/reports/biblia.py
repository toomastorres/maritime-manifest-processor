"""
biblia.py - Biblia del viaje: 'ATL 0426 - IMPO.xlsx' / 'ATL 0426 - EXPO.xlsx'.

Una hoja por puerto nuestro (ZARATE, ...) con UNA tabla: el puerto del exterior
es una columna (POL en IMPO, POD en EXPO) en lugar de un bloque por puerto, así
la planilla es más corta y se puede filtrar. Los totales de la tabla usan
SUBTOTALES: al filtrar por puerto o por tipo, el total muestra lo filtrado.

Cada B/L ocupa una fila; si trae cargas que comisionan distinto (contenedor,
Bbulk, rodado) ocupa una fila por carga (ver rules/split.py). Lo que es del B/L
entero (collect, BAF) va en la primera fila.

Columnas (IMPO):
  B/L | POL | CARGA | TIPO | CANTIDADES 20' 40' CAR VAN RORO BBULK |
  FACTURAMOS NOSOTROS: PESO DTN TOLL THC 20' THC 40' SWP |
  FLETE: USD PREPAID, USD COLLECT, EUR PREPAID, EUR COLLECT, BAF USD, BAF EUR |
  COMISIÓN: REGLA, USD | COBRANZA: TOLL, THC 20', THC 40', SWP cobrados (a mano) y PENDIENTE
EXPO: igual hasta los gastos (más ENS), FLETE: USD PREPAID (incl. ENS),
  COLLECT USD/EUR/GBP/BRL, ABROAD | COMISIÓN.

Hojas: una por puerto nuestro, RESUMEN (por puerto del exterior, fletes y
controles), TOLL (por categoría, con SUMAR.SI) y REVISIÓN.

Convenciones operativas: el TOLL prepaid muestra el puerto de origen; el
THC/SWP prepaid muestra 'Prepaid'; en EXPO, si el THC/SWP no entra en el
prepaid de Buenos Aires, se paga en Casa Matriz y se marca 'CASA MATRIZ'.
"""

from collections import OrderedDict
from dataclasses import dataclass
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import config
from ..catalog import load_catalog
from ..domain.classification import abreviar_puerto, format_cargo
from ..rules.billing import expo_money, impo_money, local_lines, total_mismatches
from ..rules.commission import (
    G_CTR, G_ROD, MODE_AGENCY, MODE_ESTIMADO, _commissionable_lines, _fmt, evaluate,
)
from ..rules.split import MIXTO, SIN_CARGA, TIPO, filas_del_bl, split_alerts
from ..validation import Alerta, revisar_bls
from ._style import add_revision_sheet, format_val, sheet_title

# ── estilos ───────────────────────────────────────────────────────────
F_TITLE = Font(name="Arial", size=16, bold=True)
F_BOLD = Font(name="Arial", size=10, bold=True)
F_DATA = Font(name="Arial", size=10)
F_GREY = Font(name="Arial", size=10, color="808080")
F_GROUP = Font(name="Arial", size=10, bold=True, italic=True)
F_TOTAL = Font(name="Arial", size=11, bold=True)
F_SECTION = Font(name="Arial", size=12, bold=True, italic=True)
F_PEND = Font(name="Arial", size=10, bold=True, color="9C0006")
F_BANNER = Font(name="Arial", size=11, bold=True, color="9C0006")

FILL_SECTION = PatternFill("solid", start_color="FFE2EFDA", end_color="FFE2EFDA")
FILL_HDR = PatternFill("solid", start_color="FFD9E1F2", end_color="FFD9E1F2")
FILL_TOTAL = PatternFill("solid", start_color="FFFCE4D6", end_color="FFFCE4D6")
FILL_PEND = PatternFill("solid", start_color="FFFFC7CE", end_color="FFFFC7CE")
FILL_INPUT = PatternFill("solid", start_color="FFFFF2CC", end_color="FFFFF2CC")

_t = Side(style="thin")
_m = Side(style="medium")
B_THIN = Border(left=_t, right=_t, top=_t, bottom=_t)
B_MED = Border(left=_m, right=_m, top=_m, bottom=_m)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
HCENTER = Alignment(horizontal="center")

MONEY = "#,##0.00"
INT = "#,##0"
TONS = "#,##0.000"
ROE_FMT = "0.0000"
MONEY_Z = "#,##0.00;-#,##0.00;"
INT_Z = "#,##0;-#,##0;"


@dataclass(frozen=True)
class Col:
    key: str
    title: str
    width: float
    fmt: Optional[str] = None
    group: Optional[str] = None
    total: bool = False          # lleva SUBTOTAL en la fila TOTAL
    center: bool = False


CANT = "CANTIDADES"
NOS = "FACTURAMOS NOSOTROS"
FLETE = "FLETE"
COMM = "COMISIÓN"
COBR = "COBRANZA (completar lo cobrado)"

_QTY = [Col("q20", "20'", 5.7, INT, CANT, True, True), Col("q40", "40'", 5.7, INT, CANT, True, True),
        Col("car", "CAR", 6.3, INT, CANT, True, True), Col("van", "VAN", 6.3, INT, CANT, True, True),
        Col("roro", "RORO", 6.7, INT, CANT, True, True), Col("bbulk", "BBULK", 7.7, INT, CANT, True, True)]

IMPO_COLS = [
    Col("bl", "B/L", 13.3), Col("puerto", "POL", 15), Col("carga", "CARGA", 32), Col("tipo", "TIPO", 13),
    *_QTY,
    Col("peso", "PESO", 9.5, TONS, NOS), Col("dtn", "DTN", 6, None, NOS, center=True),
    Col("toll", "TOLL", 11.5, MONEY, NOS, True), Col("thc20", "THC 20'", 10.5, MONEY, NOS, True),
    Col("thc40", "THC 40'", 10.5, MONEY, NOS, True), Col("swp", "SWP", 9, MONEY, NOS, True),
    Col("usd_pre", "USD PREPAID", 12.5, MONEY, FLETE, True), Col("usd_col", "USD COLLECT", 12.5, MONEY, FLETE, True),
    Col("eur_pre", "EUR PREPAID", 12.5, MONEY, FLETE, True), Col("eur_col", "EUR COLLECT", 12.5, MONEY, FLETE, True),
    Col("baf_usd", "BAF USD", 11, MONEY, FLETE, True), Col("baf_eur", "BAF EUR", 11, MONEY, FLETE, True),
    Col("regla", "REGLA", 34, None, COMM), Col("comm", "USD", 11, MONEY, COMM, True),
    Col("gap", "", 2),
    Col("c_toll", "TOLL", 11.5, MONEY, COBR, True), Col("c_thc20", "THC 20'", 10.5, MONEY, COBR, True),
    Col("c_thc40", "THC 40'", 10.5, MONEY, COBR, True), Col("c_swp", "SWP", 9, MONEY, COBR, True),
    Col("pend", "PENDIENTE", 12.5, MONEY, COBR, True),
]

EXPO_COLS = [
    Col("bl", "B/L", 13.3), Col("puerto", "POD", 15), Col("carga", "CARGA", 32), Col("tipo", "TIPO", 13),
    *_QTY,
    Col("peso", "PESO", 9.5, TONS, NOS), Col("toll", "TOLL", 11.5, MONEY, NOS, True),
    Col("thc20", "THC 20'", 10.5, MONEY, NOS, True), Col("thc40", "THC 40'", 10.5, MONEY, NOS, True),
    Col("swp", "SWP", 9, MONEY, NOS, True), Col("ens", "ENS", 8, MONEY, NOS, True),
    Col("usd_pre", "USD PREPAID (INCL ENS)", 14, MONEY, FLETE, True),
    Col("col_usd", "COLLECT USD", 12.5, MONEY, FLETE, True), Col("col_eur", "COLLECT EUR", 12.5, MONEY, FLETE, True),
    Col("col_gbp", "COLLECT GBP", 12.5, MONEY, FLETE, True), Col("col_brl", "COLLECT BRL", 12.5, MONEY, FLETE, True),
    Col("abroad", "ABROAD USD", 12.5, MONEY, FLETE, True),
    Col("regla", "REGLA", 34, None, COMM), Col("comm", "USD", 11, MONEY, COMM, True),
]

HDR_GROUP_ROW = 5
HDR_ROW = 6
FIRST_ROW = 7

PENDIENTE_REGLA = "REVISAR (ver hoja REVISIÓN)"


class _Tabla:
    """Columnas de una operación: letra por clave."""

    def __init__(self, cols: List[Col]):
        self.cols = cols
        self.idx = {c.key: i for i, c in enumerate(cols, start=1)}

    def L(self, key: str) -> str:
        return get_column_letter(self.idx[key])

    def has(self, key: str) -> bool:
        return key in self.idx


@dataclass
class _Hoja:
    title: str
    first: int
    last: int


# ───────────────────────────── API ─────────────────────────────

def generar_biblia(bls, op_type, ship_name, voyage, output_path, cfg=None, catalog=None,
                   extra_alerts=None):
    cfg = cfg or config.EngineConfig()
    catalog = catalog or load_catalog()
    op = op_type.upper()
    T = _Tabla(IMPO_COLS if op == "IMPO" else EXPO_COLS)

    resultados = {id(bl): evaluate(bl, op, cfg, catalog) for bl in bls}
    filas_por_bl = {id(bl): filas_del_bl(bl, resultados[id(bl)], op) for bl in bls}

    wb = Workbook()
    wb.remove(wb.active)
    hojas: List[_Hoja] = []
    puertos_ext: "OrderedDict[str, int]" = OrderedDict()   # puerto del exterior -> cantidad de B/L

    nuestro = (lambda b: b.port_of_discharge) if op == "IMPO" else (lambda b: b.port_of_loading)
    exterior = (lambda b: b.port_of_loading) if op == "IMPO" else (lambda b: b.port_of_discharge)

    for port, port_bls in _agrupar(bls, nuestro, ordenar=(op == "EXPO")).items():
        ws = wb.create_sheet(title=sheet_title(port))
        _encabezado(ws, T, op, ship_name, voyage, port, cfg,
                    [resultados[id(b)] for b in port_bls])
        r = FIRST_ROW
        for ext, ext_bls in _agrupar(port_bls, exterior, ordenar=(op == "EXPO")).items():
            puertos_ext[ext] = puertos_ext.get(ext, 0) + len(ext_bls)
            for bl in ext_bls:
                for f in filas_por_bl[id(bl)]:
                    if op == "IMPO":
                        _fila_impo(ws, T, r, f, port, cfg, catalog)
                    else:
                        _fila_expo(ws, T, r, f, cfg)
                    r += 1
        last = r - 1
        _fila_total(ws, T, r, FIRST_ROW, last)
        _hoja_setup(ws, T, last)
        hojas.append(_Hoja(ws.title, FIRST_ROW, max(last, FIRST_ROW)))

    if not hojas:
        wb.create_sheet(title="SIN DATOS")
    else:
        _resumen(wb, T, op, ship_name, voyage, hojas, puertos_ext, bls, resultados, cfg)
        _toll(wb, T, hojas, cfg, any(f.tipo in (MIXTO, SIN_CARGA) and f.toll_lines
                                     for fs in filas_por_bl.values() for f in fs))

    alertas = []
    for bl in bls:
        res = resultados[id(bl)]
        if not res.ok:
            detalle = " | ".join(res.reasons) + (f" — {res.suggestion}" if res.suggestion else "")
            alertas.append(Alerta(bl.bl_no, "COMISIÓN - REVISAR", detalle))
        for msg in split_alerts(filas_por_bl[id(bl)]):
            alertas.append(Alerta(bl.bl_no, "CARGA MIXTA", msg))
    alertas += list(extra_alerts or [])
    alertas += revisar_bls(bls, op, cfg, catalog=catalog)
    _imprimir(add_revision_sheet(wb, alertas))
    wb.save(output_path)
    return output_path


# ───────────────────────────── filas ─────────────────────────────

def _agrupar(bls, key, ordenar):
    out = OrderedDict()
    claves = sorted({key(b) for b in bls}) if ordenar else list(OrderedDict.fromkeys(key(b) for b in bls))
    for k in claves:
        grupo = [b for b in bls if key(b) == k]
        out[k] = sorted(grupo, key=lambda b: b.bl_no) if ordenar else grupo
    return out


def _carga(f) -> object:
    """Contenedores: la cantidad. Si no, el texto descriptivo."""
    c = f.cargo
    if c.containers and not c.vehicles and not (c.general and not f.extra_items):
        return c.n_containers
    return format_cargo(f.bl, c) or None


def _regla(f) -> str:
    res = f.res
    if not res.ok:
        return PENDIENTE_REGLA
    textos = []
    for p in f.parts:
        if p.group == G_ROD:
            if p.mode == MODE_AGENCY:
                pct_rod = f"{p.rate * 100:g}%".replace(".", ",")
                textos.append(f"{pct_rod} s/ {p.currency} {_fmt(p.base)} (AMA)")
            else:
                u = " + ".join(f"{n} {fam} x {config.PER_UNIT_TARIFF_USD[fam]:g} USD"
                               for fam, n in p.units.items() if n)
                textos.append(u + (" (estimado)" if p.mode == MODE_ESTIMADO else ""))
            continue
        if not p.base and not p.commission_usd:
            continue
        pct = f"{p.rate * 100:g}%".replace(".", ",")
        t = f"{pct} s/ {p.currency} {_fmt(p.base)}"
        if res.all_in:
            t += " (ALL IN 80%)"
        if res.brasil:
            t += " Brasil"
        if p.group == G_CTR and p.minimum_usd > p.calc_usd:
            min_u = p.minimum_usd / p.min_ctr_units if p.min_ctr_units else 0
            t = f"MÍNIMO {min_u:g} x {p.min_ctr_units} ctr (el {pct} da {_fmt(p.calc_usd)})"
        textos.append(t)
    return " + ".join(textos)


def _comm(f):
    if not f.res.ok:
        return "REVISAR"
    # Sin redondear: así el total coincide al centavo con la planilla de comisiones.
    return sum(p.commission_usd for p in f.parts)


def _put(ws, T, r, key, value, fmt=None):
    if not T.has(key) or value is None or value == "":
        return None
    c = ws.cell(row=r, column=T.idx[key], value=value)
    if fmt:
        c.number_format = fmt
    return c


def _comunes(ws, T, r, f, puerto_ext):
    bl = f.bl
    _put(ws, T, r, "bl", bl.bl_no)
    _put(ws, T, r, "puerto", puerto_ext)
    _put(ws, T, r, "carga", _carga(f))
    _put(ws, T, r, "tipo", f.tipo)
    for key, val in (("q20", f.qty_size("20")), ("q40", f.qty_size("40")), ("car", f.qty("CAR")),
                     ("van", f.qty("VAN")), ("roro", f.qty("RORO")), ("bbulk", f.qty("BBULK"))):
        ws.cell(row=r, column=T.idx[key], value=val)


def _peso(ws, T, r, f, toll_es_numero, cfg):
    n_ctr = f.qty_size("20") + f.qty_size("40")
    if f.tipo == TIPO[G_CTR] and n_ctr:
        _put(ws, T, r, "peso", n_ctr, INT)
    elif toll_es_numero and f.tipo != TIPO[G_CTR]:
        _put(ws, T, r, "peso", f"={T.L('toll')}{r}/{cfg.tonnage_factor:g}")
    elif f.tons:
        _put(ws, T, r, "peso", f.tons)


def _estilo_fila(ws, T, r, f):
    for col in T.cols:
        c = ws.cell(row=r, column=T.idx[col.key])
        if col.key == "gap":
            continue
        c.font = F_DATA
        c.border = B_THIN
        if col.fmt and not c.number_format.startswith('"'):
            if c.number_format == "General":
                c.number_format = col.fmt
        if col.center:
            c.alignment = HCENTER
        if col.group == COBR and col.key != "pend":
            c.fill = FILL_INPUT
    if not f.first:
        ws.cell(row=r, column=T.idx["bl"]).font = F_GREY
    if not f.res.ok:
        for key in ("bl", "regla", "comm"):
            c = ws.cell(row=r, column=T.idx[key])
            c.fill = FILL_PEND
            c.font = F_PEND


def _fila_impo(ws, T, r, f, port, cfg, catalog):
    bl = f.bl
    _comunes(ws, T, r, f, bl.port_of_loading)
    _put(ws, T, r, "dtn", abreviar_puerto(port))

    # TOLL: lo que facturamos nosotros; si sólo hay TOLL prepaid (pagado en
    # origen), se muestra el puerto de origen.
    toll_es_numero = False
    if f.toll_lines:
        _put(ws, T, r, "toll", f.toll)
        toll_es_numero = True
    elif f.toll_prepaid:
        _put(ws, T, r, "toll", bl.port_of_loading)
    _peso(ws, T, r, f, toll_es_numero, cfg)

    if f.container_costs:
        _put(ws, T, r, "thc20", format_val(bl.thc_20.monto, bl.thc_20.letra, item_type="THC"))
        _put(ws, T, r, "thc40", format_val(bl.thc_40.monto, bl.thc_40.letra, item_type="THC"))
        swp = bl.charge("Sweeping")
        _put(ws, T, r, "swp", format_val(swp.monto, swp.letra, item_type="SWP"))

    # Flete prepaid: lo comisionable de esta carga, por moneda.
    lineas = f.comm_lines
    if f.tipo == SIN_CARGA and not lineas:
        lineas = _commissionable_lines(bl, catalog, "IMPO")
    pre = {}
    for cl in lineas:
        if cl.letra == "P":
            pre[cl.moneda] = pre.get(cl.moneda, 0.0) + cl.total
    _put(ws, T, r, "usd_pre", _num(max(0.0, pre.get("USD", 0.0))))
    _put(ws, T, r, "eur_pre", _num(max(0.0, pre.get("EUR", 0.0))))

    if f.first:
        m = impo_money(bl)
        _put(ws, T, r, "usd_col", _num(m.collect.get("USD")))
        _put(ws, T, r, "eur_col", _num(m.collect.get("EUR")))
        _put(ws, T, r, "baf_usd", _num(m.baf.get("USD")))
        _put(ws, T, r, "baf_eur", _num(m.baf.get("EUR")))

    _put(ws, T, r, "regla", _regla(f))
    _put(ws, T, r, "comm", _comm(f))
    _put(ws, T, r, "pend",
         f"=SUM({T.L('toll')}{r}:{T.L('swp')}{r})-SUM({T.L('c_toll')}{r}:{T.L('c_swp')}{r})")
    _estilo_fila(ws, T, r, f)


def _fila_expo(ws, T, r, f, cfg):
    bl = f.bl
    _comunes(ws, T, r, f, bl.port_of_discharge)

    toll_es_numero = bool(f.toll_lines)
    if f.toll_lines:
        _put(ws, T, r, "toll", f.toll)
    _peso(ws, T, r, f, toll_es_numero, cfg)

    m = expo_money(bl)
    swp = bl.charge("Sweeping")
    gastos = ((bl.thc_20, "thc20", "thc (pol)"), (bl.thc_40, "thc40", "thc (pol)"), (swp, "swp", "sweeping"))
    casa_matriz = _pagados_afuera(bl, m, gastos)

    if f.container_costs:
        for g, key, _ in gastos:
            if not g.monto:
                continue
            if g.letra == "C":
                _put(ws, T, r, key, "Collect")
            elif key in casa_matriz:
                _put(ws, T, r, key, "CASA MATRIZ")
            else:
                _put(ws, T, r, key, g.monto)

    if f.first:
        ens = bl.charge("ENS")
        _put(ws, T, r, "ens", ens.monto if ens.monto else "***")
        refs = [key for g, key, _ in gastos if g.monto and g.letra == "P" and key not in casa_matriz]
        if m.local_prepaid or refs:
            if refs and f.container_costs:
                resta = "".join(f"-{T.L(k)}{r}" for k in refs)
                _put(ws, T, r, "usd_pre", f"={_num_txt(m.local_prepaid)}{resta}")
            else:
                _put(ws, T, r, "usd_pre", m.local_prepaid)
        for moneda, key in (("USD", "col_usd"), ("EUR", "col_eur"), ("GBP", "col_gbp"), ("BRL", "col_brl")):
            _put(ws, T, r, key, _num(m.collect.get(moneda)))
        abroad = dict(m.abroad)
        usd = abroad.pop("USD", None)
        otras = " + ".join(f"{k} {v:,.2f}" for k, v in abroad.items())
        if usd:
            c = _put(ws, T, r, "abroad", _num(usd))
            if otras and c is not None:
                c.comment = Comment(f"Además, prepaid afuera: {otras}", "Motor")
        elif otras:
            _put(ws, T, r, "abroad", otras)

    _put(ws, T, r, "regla", _regla(f))
    _put(ws, T, r, "comm", _comm(f))
    _estilo_fila(ws, T, r, f)


def _pagados_afuera(bl, m, gastos) -> set:
    """
    THC / Sweeping prepaid que se pagan en Casa Matriz y no en Buenos Aires.

    Si el manifiesto tiene un prepaid en Buenos Aires y otro afuera, se identifica
    qué líneas forman el de Buenos Aires (combinación única que suma su total).
    Si no se puede identificar: si el THC y el Sweeping no entran en el prepaid
    de Buenos Aires, se pagan en Casa Matriz.
    """
    prepagos = [(g, key, desc) for g, key, desc in gastos if g.monto and g.letra == "P"]
    if not prepagos:
        return set()
    locales = local_lines(bl, "P", "USD")
    if locales is not None:
        ids = {id(cl) for cl in locales}
        afuera = set()
        for g, key, desc in prepagos:
            propias = [cl for cl in bl.charge_lines if cl.letra == "P" and desc in cl.desc]
            if propias and not any(id(cl) in ids for cl in propias):
                afuera.add(key)
        return afuera
    descontado = sum(g.monto for g, _, _ in prepagos)
    if m.local_prepaid and m.local_prepaid - descontado < 0:
        return {key for _, key, _ in prepagos}
    return set()


def _num(v):
    return round(v, 2) if v is not None and abs(v) > 0.005 else None


def _num_txt(v: float) -> str:
    return ("%f" % v).rstrip("0").rstrip(".") or "0"


# ─────────────────────────── encabezado y total ───────────────────────────

def _encabezado(ws, T, op, ship, voyage, port, cfg, resultados):
    lugar = "DESCARGA" if op == "IMPO" else "CARGA"
    ws.cell(row=1, column=1, value=f"{op} — {ship} {voyage} — {lugar} {port}").font = F_TITLE
    ws.row_dimensions[1].height = 20.25
    ws.cell(row=2, column=1, value="ROE: EU 1,00 =").font = F_BOLD
    c = ws.cell(row=2, column=2, value=cfg.roe)
    c.font = F_BOLD
    c.number_format = ROE_FMT
    c.alignment = Alignment(horizontal="left")
    avisos = []
    if cfg.agency_bls is None:
        avisos.append("ESTIMACIÓN: sin lista de Comercial, rodados por unidad")
    pend = sum(1 for r in resultados if not r.ok)
    if pend:
        avisos.append(f"{pend} B/L CON COMISIÓN PENDIENTE DE REVISIÓN (en rojo, ver hoja REVISIÓN)")
    if avisos:
        ws.cell(row=3, column=1, value=" | ".join(avisos)).font = F_BANNER

    # Fila de grupos (combinada por grupo) y fila de títulos.
    i = 1
    while i <= len(T.cols):
        col = T.cols[i - 1]
        j = i
        while j < len(T.cols) and col.group and T.cols[j].group == col.group:
            j += 1
        if col.group:
            ws.cell(row=HDR_GROUP_ROW, column=i, value=col.group)
            for k in range(i, j + 1):
                c = ws.cell(row=HDR_GROUP_ROW, column=k)
                c.font = F_GROUP
                c.fill = FILL_SECTION
                c.border = B_THIN
                c.alignment = CENTER
            if j > i:
                ws.merge_cells(start_row=HDR_GROUP_ROW, start_column=i, end_row=HDR_GROUP_ROW, end_column=j)
        i = j + 1
    for col in T.cols:
        c = ws.cell(row=HDR_ROW, column=T.idx[col.key], value=col.title or None)
        ws.column_dimensions[get_column_letter(T.idx[col.key])].width = col.width
        if col.key == "gap":
            continue
        c.font = F_BOLD
        c.fill = FILL_INPUT if (col.group == COBR and col.key != "pend") else FILL_HDR
        c.border = B_THIN
        c.alignment = CENTER
    ws.row_dimensions[HDR_ROW].height = 27


def _fila_total(ws, T, r, first, last):
    ws.cell(row=r, column=1, value="TOTAL")
    for col in T.cols:
        if col.key == "gap":
            continue
        c = ws.cell(row=r, column=T.idx[col.key])
        if col.total and last >= first:
            L = T.L(col.key)
            c.value = f"=SUBTOTAL(9,{L}{first}:{L}{last})"
            c.number_format = col.fmt or MONEY
        c.font = F_TOTAL
        c.fill = FILL_TOTAL
        c.border = B_MED
        if col.center:
            c.alignment = HCENTER
    ws.row_dimensions[r].height = 16.5


def _hoja_setup(ws, T, last):
    ws.sheet_view.showZeros = False
    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=T.idx["tipo"] + 1)
    ultima = get_column_letter(len(T.cols))
    ws.auto_filter.ref = f"A{HDR_ROW}:{ultima}{max(last, HDR_ROW + 1)}"
    _imprimir(ws)
    ws.print_title_rows = f"{HDR_GROUP_ROW}:{HDR_ROW}"


# ─────────────────────────── fórmulas entre hojas ───────────────────────────

def _rng(h: _Hoja, L: str) -> str:
    return f"'{h.title}'!${L}${h.first}:${L}${h.last}"


def _sum(T, hojas, key):
    return "=" + "+".join(f"SUM({_rng(h, T.L(key))})" for h in hojas)


def _sumifs(T, hojas, key, conds, eq=True):
    """SUMAR.SI.CONJUNTO sumado en todas las hojas. conds: [(clave, criterio)]."""
    partes = []
    for h in hojas:
        args = ",".join(f"{_rng(h, T.L(k))},{crit}" for k, crit in conds)
        partes.append(f"SUMIFS({_rng(h, T.L(key))},{args})")
    return ("=" if eq else "") + "+".join(partes)


# ─────────────────────────── RESUMEN ───────────────────────────

def _resumen(wb, T, op, ship, voyage, hojas, puertos_ext, bls, resultados, cfg):
    ws = wb.create_sheet(title="RESUMEN")
    ws.cell(row=1, column=1, value=f"RESUMEN {op} — {ship} {voyage}").font = F_TITLE
    ws.row_dimensions[1].height = 20.25

    numericas = [c for c in T.cols if c.total and c.group != COBR]
    cols = ["puerto", "bls"] + [c.key for c in numericas]
    ws.column_dimensions["A"].width = 30
    for i in range(2, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12.5

    def section(r, text, ancho):
        c = ws.cell(row=r, column=1, value=text)
        c.font = F_SECTION
        for k in range(1, ancho + 1):
            ws.cell(row=r, column=k).fill = FILL_SECTION
            ws.cell(row=r, column=k).border = B_THIN
        return r + 1

    # 1. Por puerto del exterior.
    r = section(3, "POR PUERTO DE ORIGEN" if op == "IMPO" else "POR PUERTO DE DESTINO", len(cols))
    titulos = ["PUERTO", "B/L"] + [(c.title if c.group != COMM else "COMISIÓN USD") for c in numericas]
    for i, t in enumerate(titulos, start=1):
        c = ws.cell(row=r, column=i, value=t)
        c.font = F_BOLD
        c.fill = FILL_HDR
        c.border = B_THIN
        c.alignment = CENTER
    ws.row_dimensions[r].height = 27
    r += 1
    primera = r
    for puerto, n in puertos_ext.items():
        ws.cell(row=r, column=1, value=puerto)
        ws.cell(row=r, column=2, value=n)
        for i, col in enumerate(numericas, start=3):
            c = ws.cell(row=r, column=i, value=_sumifs(T, hojas, col.key, [("puerto", f"$A{r}")]))
            c.number_format = INT_Z if col.fmt == INT else MONEY_Z
        for i in range(1, len(cols) + 1):
            ws.cell(row=r, column=i).font = F_DATA
            ws.cell(row=r, column=i).border = B_THIN
        r += 1
    ultima = r - 1
    ws.cell(row=r, column=1, value="TOTAL")
    for i in range(2, len(cols) + 1):
        L = get_column_letter(i)
        c = ws.cell(row=r, column=i, value=f"=SUM({L}{primera}:{L}{ultima})")
        c.number_format = INT if i == 2 else (INT_Z if numericas[i - 3].fmt == INT else MONEY_Z)
    for i in range(1, len(cols) + 1):
        c = ws.cell(row=r, column=i)
        c.font = F_TOTAL
        c.fill = FILL_TOTAL
        c.border = B_MED
    fila_total = r
    col_de = {c.key: get_column_letter(i) for i, c in enumerate(numericas, start=3)}
    r += 2

    def linea(r, label, formula=None, fmt=MONEY, bold=False, nota=None, fill=None):
        a = ws.cell(row=r, column=1, value=label)
        a.font = F_BOLD if bold else F_DATA
        if formula is not None:
            b = ws.cell(row=r, column=2, value=formula)
            b.font = F_BOLD if bold else F_DATA
            b.number_format = fmt
            b.border = B_THIN
            if fill:
                b.fill = fill
        if nota:
            ws.cell(row=r, column=3, value=nota).font = Font(name="Arial", size=9, italic=True, color="595959")
        return r + 1

    def ok_if(cond):
        return f'=IF({cond},"OK","REVISAR")'

    if op == "EXPO":
        r = section(r, "FLETES (para el INFOFRTS)", 6)
        r = linea(r, "FREIGHTS PREPAID (USD, incl. ENS)", f"={col_de['usd_pre']}{fila_total}")
        for moneda, key in (("USD", "col_usd"), ("EUR", "col_eur"), ("GBP", "col_gbp"), ("BRL", "col_brl")):
            r = linea(r, f"FREIGHTS COLLECT {moneda}", f"={col_de[key]}{fila_total}")
        r = linea(r, "FREIGHTS ABROAD (USD)", f"={col_de['abroad']}{fila_total}")
        otros = {}
        for bl in bls:
            for k, v in expo_money(bl).abroad.items():
                if k != "USD":
                    otros[k] = otros.get(k, 0.0) + v
        for k, v in sorted(otros.items()):
            r = linea(r, f"FREIGHTS ABROAD ({k})", round(v, 2),
                      nota="no suma en la columna ABROAD USD: figura como texto o comentario")
        r += 1

    if op == "IMPO":
        for moneda, key in (("USD", "usd_col"), ("EUR", "eur_col")):
            r = section(r, f"FLETES COLLECT {moneda}", 3)
            c_col = r
            r = linea(r, "COLLECT", f"={col_de[key]}{fila_total}")
            c_af = r
            r = linea(r, "LO PAGAN AFUERA", 0, fill=FILL_INPUT, nota="completar a mano si corresponde")
            r = linea(r, "TOTAL A COBRAR", f"=B{c_col}-B{c_af}", bold=True)
            r += 1

        r = section(r, "COBRANZA (lo que facturamos nosotros)", 4)
        for i, t in enumerate(("", "FACTURADO", "COBRADO", "PENDIENTE"), start=1):
            c = ws.cell(row=r, column=i, value=t or None)
            c.font = F_BOLD
            if t:
                c.fill = FILL_HDR
                c.border = B_THIN
                c.alignment = CENTER
        r += 1
        for label, key, ckey in (("TOLL", "toll", "c_toll"), ("THC 20'", "thc20", "c_thc20"),
                                 ("THC 40'", "thc40", "c_thc40"), ("SWEEPING", "swp", "c_swp")):
            ws.cell(row=r, column=1, value=label).font = F_DATA
            for i, f in ((2, _sum(T, hojas, key)), (3, _sum(T, hojas, ckey)), (4, f"=B{r}-C{r}")):
                c = ws.cell(row=r, column=i, value=f)
                c.number_format = MONEY
                c.font = F_DATA
                c.border = B_THIN
            r += 1
        r += 1

    # 4. Controles.
    r = section(r, "CONTROLES", 4)
    if op == "IMPO":
        p20 = sum(round(bl.thc_20.monto / cfg.tariffs.thc_20) for bl in bls if bl.thc_20.letra == "P")
        p40 = sum(round(bl.thc_40.monto / cfg.tariffs.thc_40) for bl in bls if bl.thc_40.letra == "P")
        for i, t in enumerate(("CONTENEDORES", "20'", "40'", "TOTAL"), start=1):
            c = ws.cell(row=r, column=i, value=t)
            c.font = F_BOLD
            c.fill = FILL_HDR
            c.border = B_THIN
        r += 1
        filas = (("Según manifiesto", _sum(T, hojas, "q20"), _sum(T, hojas, "q40"), INT),
                 (f"THC cobrado acá (÷{cfg.tariffs.thc_20:g} / ÷{cfg.tariffs.thc_40:g})",
                  f"={col_de['thc20']}{fila_total}/{cfg.tariffs.thc_20:g}",
                  f"={col_de['thc40']}{fila_total}/{cfg.tariffs.thc_40:g}", "0.00"),
                 ("THC prepaid (se paga en origen)", p20, p40, INT))
        ref = {}
        for label, a, b, fmt in filas:
            ws.cell(row=r, column=1, value=label).font = F_DATA
            for i, v in ((2, a), (3, b), (4, f"=B{r}+C{r}")):
                c = ws.cell(row=r, column=i, value=v)
                c.number_format = fmt
                c.font = F_DATA
                c.border = B_THIN
            ref[label[:5]] = r
            r += 1
        ws.cell(row=r, column=1, value=f"Según TOLL de contenedores (÷{cfg.tariffs.sc_container:g})").font = F_DATA
        c = ws.cell(row=r, column=4, value="='TOLL'!C11")
        c.number_format = "0.00"
        c.border = B_THIN
        t_toll = r
        r += 1
        man, aca, pre = ref["Según"], ref["THC c"], ref["THC p"]
        r = linea(r, "Manifiesto = THC cobrado acá + THC prepaid",
                  ok_if(f"ABS(D{man}-D{aca}-D{pre})<0.01"), fmt="General", bold=True)
        r = linea(r, "Manifiesto = contenedores según TOLL",
                  ok_if(f"ABS(D{man}-D{t_toll})<0.01"), fmt="General", bold=True,
                  nota="los contenedores con TOLL prepaid no figuran en el TOLL")
        r += 1

    mal_total = sum(1 for bl in bls if total_mismatches(bl))
    r = linea(r, "B/L cuyas líneas de cargo no suman sus Total[...]", mal_total, fmt=INT, bold=True,
              nota="OK" if not mal_total else "ver hoja REVISIÓN (TOTAL MANIFIESTO)")
    r = linea(r, "TOLL por categoría = TOLL de la tabla", ok_if(f"ABS('TOLL'!D7-{col_de['toll']}{fila_total})<0.01"),
              fmt="General", bold=True)
    pend = sum(1 for res in resultados.values() if not res.ok)
    r = linea(r, "Comisión calculada (USD)", f"={col_de['comm']}{fila_total}", bold=True,
              nota="igual al total de la planilla de comisiones, sin el Control Equipment Fee")
    r = linea(r, "B/L con comisión pendiente de revisión", pend, fmt=INT, bold=True,
              nota="OK" if not pend else "en rojo en la tabla; ver hoja REVISIÓN")
    ws.freeze_panes = "B5"
    _imprimir(ws)


# ─────────────────────────── TOLL ───────────────────────────

def _toll(wb, T, hojas, cfg, hay_sin_clasificar):
    """
    Hoja TOLL con SUMIFS sobre la columna TIPO:
      CONTAINERS = TOLL de las filas CONTENEDOR
      PROJECTS   = TOLL de las filas BBULK + RODADO con RoRo
      CARS       = TOLL de las filas RODADO sin RoRo (autos y vans)
    """
    ws = wb.create_sheet(title="TOLL")
    for col, w in (("A", 3), ("B", 30), ("C", 15), ("D", 15), ("E", 15), ("F", 17)):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showZeros = False

    def hdr(r, values):
        for col, v in values.items():
            c = ws[f"{col}{r}"]
            c.value = v
            c.font = F_BOLD
            c.fill = FILL_HDR
            c.border = B_MED
            c.alignment = CENTER

    cats = {"C": "CONTAINERS", "D": "PROJECTS", "E": "CARS"}
    if hay_sin_clasificar:
        cats["F"] = "SIN CLASIFICAR"
    hdr(3, {"B": "///////////////////", **cats})
    ws["B4"] = "TOLL"
    s = lambda conds, eq=False: _sumifs(T, hojas, "toll", conds, eq)
    ws["C4"] = s([("tipo", '"CONTENEDOR"')], True)
    rodado_roro = s([("tipo", '"RODADO"'), ("roro", '">0"')])
    ws["D4"] = "=" + s([("tipo", '"BBULK"')]) + "+" + rodado_roro
    ws["E4"] = "=" + s([("tipo", '"RODADO"')]) + "-(" + rodado_roro + ")"
    if hay_sin_clasificar:
        ws["F4"] = "=" + s([("tipo", f'"{MIXTO}"')]) + "+" + s([("tipo", f'"{SIN_CARGA}"')])
        ws["F5"] = "B/L con varias cargas cuyo TOLL no se pudo separar: clasificar a mano"
        ws["F5"].font = Font(name="Arial", size=9, italic=True, color="9C0006")
    for col in ["B"] + list(cats):
        c = ws[f"{col}4"]
        c.font = F_DATA
        c.border = B_THIN
        if col != "B":
            c.number_format = MONEY

    ws["C7"] = "TOTAL:"
    ws["D7"] = "=" + "+".join(f"{c}4" for c in cats)
    ws["C7"].font = F_TOTAL
    ws["D7"].font = F_TOTAL
    ws["D7"].number_format = MONEY
    ws["D7"].fill = FILL_TOTAL
    ws["D7"].border = B_MED

    hdr(10, {"B": "///////////////////", "C": "CONTAINERS", "D": "PROJECTS", "E": "RORO", "F": "CARS"})
    ws["B11"] = "AMOUNT"
    ws["C11"] = f"=C4/{cfg.tariffs.sc_container:g}"
    con_toll = [("toll", '">0"')]
    ws["D11"] = _sumifs(T, hojas, "bbulk", [("tipo", '"BBULK"')] + con_toll)
    ws["E11"] = _sumifs(T, hojas, "roro", con_toll)
    ws["F11"] = "=" + _sumifs(T, hojas, "car", con_toll, False) + "+" + _sumifs(T, hojas, "van", con_toll, False)
    ws["B12"] = "Según manifiesto (con TOLL acá)"
    ws["C12"] = "=" + _sumifs(T, hojas, "q20", [("tipo", '"CONTENEDOR"')] + con_toll, False) + "+" + \
        _sumifs(T, hojas, "q40", [("tipo", '"CONTENEDOR"')] + con_toll, False)
    ws["B13"] = "CONTROL"
    ws["C13"] = '=IF(ABS(C11-C12)<0.01,"OK","REVISAR")'
    for r in (11, 12, 13):
        for col in "BCDEF":
            c = ws[f"{col}{r}"]
            c.font = F_BOLD if r == 13 or col == "B" else F_DATA
            if c.value is not None:
                c.border = B_THIN
                if col != "B" and r != 13:
                    c.number_format = "#,##0.00" if col == "C" and r == 11 else INT
    sc_txt = f"{cfg.tariffs.sc_container:g}".replace(".", ",")
    for i, texto in enumerate((f"AMOUNT: contenedores = TOLL ÷ {sc_txt}; PROJECTS = bultos Bbulk;",
                               "RORO y CARS = unidades. Sólo las cargas que pagan el TOLL acá."), start=15):
        nota = ws[f"B{i}"]
        nota.value = texto
        nota.font = Font(name="Arial", size=9, italic=True, color="595959")
    _imprimir(ws)


def _imprimir(ws):
    """A4 apaisado, todo el ancho en una página."""
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
